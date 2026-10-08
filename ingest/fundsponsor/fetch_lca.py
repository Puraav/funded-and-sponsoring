"""Download the DOL LCA disclosure files into data/raw/lca/lca_all.parquet.

Each quarterly Excel file is converted to parquet once (only the columns the models need),
then the quarters are stacked and de-duplicated on CASE_NUMBER, keeping the latest file's row.
"""

from __future__ import annotations

import argparse
import re
from concurrent.futures import ProcessPoolExecutor
from datetime import date, datetime
from pathlib import Path

import duckdb
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from . import config
from .http import download

STRING_COLUMNS = [
    "CASE_NUMBER", "CASE_STATUS", "VISA_CLASS", "JOB_TITLE", "SOC_CODE", "SOC_TITLE",
    "FULL_TIME_POSITION", "EMPLOYER_NAME", "TRADE_NAME_DBA", "EMPLOYER_ADDRESS1",
    "EMPLOYER_CITY", "EMPLOYER_STATE", "EMPLOYER_POSTAL_CODE", "NAICS_CODE", "WORKSITE_CITY",
    "WORKSITE_STATE", "WORKSITE_POSTAL_CODE", "WAGE_UNIT_OF_PAY", "PW_UNIT_OF_PAY",
    "PW_WAGE_LEVEL",
]  # fmt: skip
DATE_COLUMNS = ["RECEIVED_DATE", "DECISION_DATE", "BEGIN_DATE"]
INT_COLUMNS = [
    "TOTAL_WORKER_POSITIONS", "NEW_EMPLOYMENT", "CONTINUED_EMPLOYMENT", "CHANGE_EMPLOYER",
]  # fmt: skip
FLOAT_COLUMNS = ["WAGE_RATE_OF_PAY_FROM", "WAGE_RATE_OF_PAY_TO", "PREVAILING_WAGE"]

# Output column order, as listed in spec 03.
KEEP_COLUMNS = [
    "CASE_NUMBER", "CASE_STATUS", "RECEIVED_DATE", "DECISION_DATE", "VISA_CLASS", "JOB_TITLE",
    "SOC_CODE", "SOC_TITLE", "FULL_TIME_POSITION", "BEGIN_DATE", "TOTAL_WORKER_POSITIONS",
    "NEW_EMPLOYMENT", "CONTINUED_EMPLOYMENT", "CHANGE_EMPLOYER", "EMPLOYER_NAME",
    "TRADE_NAME_DBA", "EMPLOYER_ADDRESS1", "EMPLOYER_CITY", "EMPLOYER_STATE",
    "EMPLOYER_POSTAL_CODE", "NAICS_CODE", "WORKSITE_CITY", "WORKSITE_STATE",
    "WORKSITE_POSTAL_CODE", "WAGE_RATE_OF_PAY_FROM", "WAGE_RATE_OF_PAY_TO", "WAGE_UNIT_OF_PAY",
    "PREVAILING_WAGE", "PW_UNIT_OF_PAY", "PW_WAGE_LEVEL",
]  # fmt: skip

# Older column names → the names above, in case a file uses them.
ALIASES = {
    "EMPLOYER_BUSINESS_DBA": "TRADE_NAME_DBA",
    "PERIOD_OF_EMPLOYMENT_START_DATE": "BEGIN_DATE",
    "WAGE_RATE_OF_PAY_FROM_1": "WAGE_RATE_OF_PAY_FROM",
    "WAGE_RATE_OF_PAY_TO_1": "WAGE_RATE_OF_PAY_TO",
    "WAGE_UNIT_OF_PAY_1": "WAGE_UNIT_OF_PAY",
    "PREVAILING_WAGE_1": "PREVAILING_WAGE",
    "PW_UNIT_OF_PAY_1": "PW_UNIT_OF_PAY",
    "PW_WAGE_LEVEL_1": "PW_WAGE_LEVEL",
    "WORKSITE_CITY_1": "WORKSITE_CITY",
    "WORKSITE_STATE_1": "WORKSITE_STATE",
    "WORKSITE_POSTAL_CODE_1": "WORKSITE_POSTAL_CODE",
}

SCHEMA = pa.schema(
    [
        pa.field(
            name,
            pa.date32() if name in DATE_COLUMNS
            else pa.int64() if name in INT_COLUMNS
            else pa.float64() if name in FLOAT_COLUMNS
            else pa.string(),
        )
        for name in KEEP_COLUMNS
    ]
    + [
        pa.field("fiscal_year", pa.int32()),
        pa.field("fiscal_quarter", pa.int32()),
        pa.field("source_file", pa.string()),
    ]
)  # fmt: skip

_EXCEL_TEXT = re.compile(r'^="(.*)"$', re.DOTALL)


def clean_text(value) -> str | None:
    """Excel cell → trimmed text. DOL wraps many text cells as ="07310" to protect zeros."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    wrapped = _EXCEL_TEXT.match(text)
    if wrapped:
        text = wrapped.group(1).strip()
    return text or None


def clean_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def clean_number(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    text = (clean_text(value) or "").replace(",", "").replace("$", "")
    try:
        return float(text)
    except ValueError:
        return None


def clean_int(value) -> int | None:
    number = clean_number(value)
    return None if number is None else int(number)


def _cleaner(column: str):
    if column in DATE_COLUMNS:
        return clean_date
    if column in INT_COLUMNS:
        return clean_int
    if column in FLOAT_COLUMNS:
        return clean_number
    return clean_text


def read_header(xlsx: Path) -> list[str]:
    """The real header row of a DOL file, upper-cased."""
    workbook = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        first = next(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()
    return [str(c).strip().upper() if c is not None else "" for c in first]


def convert(xlsx: Path, dest: Path, fiscal_year: int, fiscal_quarter: int) -> int:
    """Convert one Excel file to parquet, keeping KEEP_COLUMNS. Returns the row count."""
    workbook = openpyxl.load_workbook(xlsx, read_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = [str(c).strip().upper() if c is not None else "" for c in next(rows)]
    header = [ALIASES.get(name, name) for name in header]
    missing = [name for name in KEEP_COLUMNS if name not in header]
    if missing:
        raise KeyError(f"{xlsx.name} is missing {missing}. Real header: {header}")
    picks = [(name, header.index(name), _cleaner(name)) for name in KEEP_COLUMNS]

    dest.parent.mkdir(parents=True, exist_ok=True)
    part = dest.with_suffix(".parquet.part")
    total = 0
    with pq.ParquetWriter(part, SCHEMA) as writer:
        batch: dict[str, list] = {name: [] for name in KEEP_COLUMNS}

        def flush() -> None:
            size = len(batch["CASE_NUMBER"])
            if not size:
                return
            table = pa.table(
                {
                    **batch,
                    "fiscal_year": [fiscal_year] * size,
                    "fiscal_quarter": [fiscal_quarter] * size,
                    "source_file": [xlsx.name] * size,
                },
                schema=SCHEMA,
            )
            writer.write_table(table)
            for values in batch.values():
                values.clear()

        for row in rows:
            if row[picks[0][1]] is None:  # blank trailing rows
                continue
            for name, index, clean in picks:
                batch[name].append(clean(row[index]) if index < len(row) else None)
            total += 1
            if total % 50_000 == 0:
                flush()
        flush()
    workbook.close()
    part.replace(dest)
    return total


def _convert_job(job: tuple[Path, Path, int, int]) -> tuple[str, int]:
    xlsx, dest, fiscal_year, fiscal_quarter = job
    return xlsx.name, convert(xlsx, dest, fiscal_year, fiscal_quarter)


def combine(parts: list[Path], dest: Path) -> None:
    """Stack the quarterly parquet files, one row per CASE_NUMBER (latest file wins)."""
    files = ", ".join(f"'{p.as_posix()}'" for p in parts)
    dest.parent.mkdir(parents=True, exist_ok=True)
    duckdb.sql(
        f"""
        copy (
            select * exclude (latest)
            from (
                select *, row_number() over (
                    partition by CASE_NUMBER
                    order by fiscal_year desc, fiscal_quarter desc
                ) as latest
                from read_parquet([{files}])
            )
            where latest = 1
            order by fiscal_year, fiscal_quarter, CASE_NUMBER
        ) to '{dest.as_posix()}' (format parquet)
        """
    )


def print_summary(parts: list[Path], dest: Path) -> None:
    files = ", ".join(f"'{p.as_posix()}'" for p in parts)
    before = duckdb.sql(f"select count(*) from read_parquet([{files}])").fetchone()[0]
    after, first, last, h1b = duckdb.sql(
        f"""select count(*), min(RECEIVED_DATE), max(RECEIVED_DATE),
                   avg((VISA_CLASS = 'H-1B')::int)
            from '{dest.as_posix()}'"""
    ).fetchone()
    print(f"\nRows read: {before:,}; after de-duplicating on CASE_NUMBER: {after:,}")
    for year, rows, low, high in duckdb.sql(
        f"""select fiscal_year, count(*), min(DECISION_DATE), max(DECISION_DATE)
            from '{dest.as_posix()}' group by 1 order by 1"""
    ).fetchall():
        print(f"  FY{year}: {rows:,} rows (decisions {low} to {high})")
    print(f"Received dates: {first} to {last}")
    print(f"Share with VISA_CLASS == 'H-1B': {h1b:.1%}")
    print(f"Saved {dest}")


def run(jobs: list[tuple[Path, Path, int, int]], dest: Path, *, workers: int = 1) -> None:
    """Convert every (xlsx, parquet, fy, quarter) job that is not done yet, then combine."""
    todo = [job for job in jobs if not job[1].exists()]
    if workers > 1 and len(todo) > 1:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            results = pool.map(_convert_job, todo)
            for name, rows in results:
                print(f"  converted {name}: {rows:,} rows")
    else:
        for job in todo:
            name, rows = _convert_job(job)
            print(f"  converted {name}: {rows:,} rows")
    combine([job[1] for job in jobs], dest)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--force", action="store_true", help="download and convert again")
    parser.add_argument("--workers", type=int, default=4, help="files converted in parallel")
    args = parser.parse_args()

    jobs = []
    for fiscal_year, quarter, url in config.LCA_FILES:
        xlsx = config.LCA_DIR / "xlsx" / url.rsplit("/", 1)[-1]
        print(f"FY{fiscal_year} Q{quarter}: {xlsx.name}")
        found = download(
            url, xlsx, user_agent=config.PUBLIC_USER_AGENT, force=args.force, progress=True
        )
        if not found:
            raise SystemExit(
                f"404 for {url}\nOpen {config.LCA_PAGE}, find the current LCA disclosure "
                "links and update LCA_FILES in config.py."
            )
        parquet = config.LCA_DIR / "parts" / f"{xlsx.stem}.parquet"
        if args.force:
            parquet.unlink(missing_ok=True)
        jobs.append((xlsx, parquet, fiscal_year, quarter))

    print(f"\nReal header of {jobs[0][0].name}:\n{', '.join(read_header(jobs[0][0]))}\n")
    run(jobs, config.LCA_PARQUET, workers=args.workers)
    print_summary([job[1] for job in jobs], config.LCA_PARQUET)


if __name__ == "__main__":
    main()
